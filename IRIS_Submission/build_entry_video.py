"""Build the IRIS entry fields and 90-second video handout.

The abstract is read from IRIS_Abstract.txt so the entry document and the
LaTeX research documents use one authoritative submission text.
"""

from pathlib import Path
import re

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
TITLE = "Physical-Only Proprioceptive Recognition of Fifteen Classes with a LEAP Hand"
CATEGORY = "Robotics and Intelligent Machines (ROBO)"
GUIDE = "Harish J. Palanthandalam-Madapusi"
GRADE = "[grade]"
ABSTRACT = (ROOT / "IRIS_Abstract.txt").read_text(encoding="utf-8").strip()

VIDEO_BLOCKS = [
    (
        "0–13 seconds",
        "Student on camera with the robotic hand visible",
        f"I'm G. Sai Pranav, a grade {GRADE} student. My project asks a simple question: can a robotic hand tell what it is holding from where its fingers stop?",
    ),
    (
        "13–34 seconds",
        "Show three short grip clips; overlay 16 + 16 + 16 = 48 angles",
        "I programmed a four-finger LEAP Hand to make three controlled grips. Motor current signals when each finger meets resistance. The recognition model receives only the final joint angles: 48 numbers per trial, with no camera image or tactile sensor array.",
    ),
    (
        "34–55 seconds",
        "Show the objects, then the empty hand; display 14 objects + empty hand = 15 labels",
        "I recorded five physical trials for each of 14 household objects and an empty-hand condition. Across the 75 trials, a small neural network's top label was correct 70 times in within-dataset validation.",
    ),
    (
        "55–73 seconds",
        "Show 70/75 and 65/67 accepted, with eight uncertain",
        "A confidence rule accepted 67 known-class decisions and got 65 right. It marked eight trials uncertain. A simpler nearest-trial comparison got 69 of 75 right in a separate evaluation.",
    ),
    (
        "73–90 seconds",
        "Return to camera; show 20 labels: 85/100 and a brief next-test card",
        "I extended the same dataset to 20 labels, with 85 of 100 raw predictions correct. These results show useful information in finger positions. My next step is testing new sessions and new copies of the objects.",
    ),
]


def word_count(value: str) -> int:
    return len(re.findall(r"\b[\w]+(?:[.-][\w]+)*\b", value))


def style_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_height = Inches(11.69)
    section.page_width = Inches(8.27)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)

    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.08

    for name, size in (("Title", 17), ("Heading 1", 11.5)):
        style = doc.styles[name]
        style.font.name = "Aptos Display" if name == "Title" else "Aptos"
        style.font.size = Pt(size)
        style.font.bold = name != "Title"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(10 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(5)
    doc.styles["Title"].paragraph_format.keep_with_next = True
    title_ppr = doc.styles["Title"]._element.get_or_add_pPr()
    for border in title_ppr.findall(qn("w:pBdr")):
        title_ppr.remove(border)


def add_blocks(doc: Document, body: str) -> None:
    for block in body.split("\n\n"):
        doc.add_paragraph(block)


def main() -> None:
    doc = Document()
    style_document(doc)
    doc.add_paragraph("IRIS Project Entry and Video Script", style="Title")
    doc.add_paragraph("Project title", style="Heading 1")
    doc.add_paragraph(TITLE)
    doc.add_paragraph("Category", style="Heading 1")
    doc.add_paragraph(CATEGORY)
    doc.add_paragraph("Single project guide", style="Heading 1")
    doc.add_paragraph(GUIDE)
    doc.add_paragraph("Project abstract", style="Heading 1")
    add_blocks(doc, ABSTRACT)

    doc.add_page_break()
    doc.add_paragraph("90 Second Project Video Script", style="Title")
    doc.add_paragraph("Only the student researcher appears on camera. Fill in the grade before recording.")
    for interval, visual, spoken in VIDEO_BLOCKS:
        doc.add_paragraph(interval, style="Heading 1")
        visual_paragraph = doc.add_paragraph()
        visual_paragraph.add_run("Visual  ").bold = True
        visual_paragraph.add_run(visual)
        spoken_paragraph = doc.add_paragraph()
        spoken_paragraph.add_run("Spoken  ").bold = True
        spoken_paragraph.add_run(spoken)
    doc.add_paragraph("Rehearse at a natural pace and keep the recorded video within 90 seconds.")
    doc.save(ROOT / "IRIS_Project_Entry_and_Video.docx")
    print("Abstract words:", word_count(ABSTRACT))
    print("Video spoken words:", sum(word_count(spoken) for _, _, spoken in VIDEO_BLOCKS))


if __name__ == "__main__":
    main()
