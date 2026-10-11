from pathlib import Path
import re
import shutil

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE_PAPER = ROOT.parent / "Manuscript" / "Manuscript.pdf"
TITLE = "Physical-Only Proprioceptive Recognition of Fifteen Classes with a LEAP Hand"
CATEGORY = "Robotics and Intelligent Machines (ROBO)"
STUDENT = "G. Sai Pranav"
GRADE = "[grade]"
GUIDE = "Harish J. Palanthandalam-Madapusi"

ABSTRACT = """Can a robotic hand recognize an object using only its own joint positions after grasping it? This project investigates that question with a physical, four-finger LEAP Hand. The hand performed the same sequence of three controlled grips on 14 household objects and in an empty-hand condition. Motor-current thresholds stopped a finger when it met resistance, but motor current was used only to control closure. Recognition used the 16 final joint angles from each grip, giving 48 angle measurements per physical trial; no camera image or tactile-array measurement entered the classifier.

Five physical trials were recorded for each of 15 labels, for a total of 75 trials. A neural network with one 24-unit hidden layer was trained on the angle vectors and evaluated in three-fold within-dataset validation. Its highest-probability label was correct for 70 of 75 held-out trials (93.3%). A decision rule accepted 67 trials as known classes; 65 of those accepted labels were correct (97.0%). The remaining eight trials were marked uncertain, and none was marked unknown. A nearest-trial comparison classified 69 of 75 trials correctly under a separate leave-one-trial-out evaluation.

These results show that the final configurations reached by repeatable physical grasps contain information that distinguishes many enrolled items under the recorded conditions. They are internal, post-selection estimates: the 15-class set was chosen after preliminary analysis, and only five trials were available per label. The study does not establish performance on a later recording session, another physical instance of an object, or a genuinely unseen class. Those tests are the next step toward reliable object recognition without external vision or dedicated tactile arrays."""

INTRO = """A dexterous robotic hand may learn about an object through the positions its joints reach when the fingers close around it. Prior studies have used joint information to infer object shape or identity, yet the value of a simple, repeatable sequence of final grasp angles on a physical LEAP Hand remains worth testing. This project asks whether those angles can distinguish common household items, including an empty-hand condition, without supplying images or tactile-array data to the recognition model. The objective was to collect physical trials under a consistent grip protocol, train an angle-only classifier, and compare its internal validation result with a nearest-trial method. The study also examined when a prediction should be withheld as uncertain. Generalization to new sessions and new physical exemplars was left for later work."""

INNOVATION = """The contribution is a physical-hand recognition pipeline that turns three controlled, sequential grasps into a compact 48-angle object signature. Motor current stops each finger when resistance is detected, while only final joint angles reach the classifier. This keeps contact control separate from recognition input. The project combines a neural classifier, a nearest-trial comparison, an uncertainty decision rule, and resumable physical-data collection. The novelty lies in this specific implemented protocol and evaluation, rather than in a claim that proprioceptive object recognition itself is new."""

METHOD = """A four-finger LEAP Hand with 16 actuated joints ran the same three-grip program for every trial. Each item was placed near a consistent point relative to the palm. During each grip, joints advanced toward configured targets. When a moving joint's signed motor current exceeded its threshold, all four joints of that finger stopped; the other fingers continued until resistance or an angle target stopped them. The final 16 joint angles were read after each grip and concatenated in grip order. Motor currents were logged but excluded from the recognition features.

The dataset comprised five physical trials for each of 14 household objects plus an empty-hand class, giving 75 labeled 48-angle vectors. Each angle was normalized using its grip-specific start and target values. The neural classifier had 48 inputs, one 24-unit tanh hidden layer, and 15 softmax outputs. Training used full-batch Adam; noisy angle copies were added only to training folds. Three-fold validation held out real physical trials from every class. A calibrated decision rule returned a class only when confidence and distance checks passed, otherwise marking it uncertain or unknown. A nearest-trial method provided a separate leave-one-trial-out comparison. The selected class set was not fixed before preliminary analysis."""

RESULTS = """The neural model's top label was correct for 70 of 75 held-out physical trials (93.3%). It accepted 67 trials as known classes, correctly labeling 65 of them (97.0% of accepted decisions); eight were uncertain and none was marked unknown. The nearest-trial method correctly labeled 69 of 75 trials (92.0%) in its different leave-one-trial-out evaluation. These figures support the central hypothesis that final grasp angles carry useful class information for the enrolled objects. They cannot be treated as an independent test: class selection followed preliminary analysis, thresholds were calibrated on the validation folds, and each label had only five physical trials. New sessions, new object instances, unseen classes, and the benefit of repeat grasps require separate evaluation."""

ACK_REFS = """I thank Harish J. Palanthandalam-Madapusi, the project's single guide, for research guidance and technical feedback. Key references are Shaw, Agarwal and Pathak (2023), LEAP Hand, arXiv:2309.06440; Vásquez and Perdereau (2017), Robotics and Autonomous Systems 98, 135-146; Sintov and Meir (2023), International Journal of Robotics Research 42, 537-561; and Kutsuzawa et al. (2024), Frontiers in Neurorobotics 18, 1466630. The complete references and technical details appear in the research paper."""

VIDEO_BLOCKS = [
    ("0-14 seconds", "On camera with the LEAP Hand visible", "I'm G. Sai Pranav, a grade " + GRADE + " student. I studied whether a robotic hand can recognize an object from how its fingers close around it."),
    ("14-37 seconds", "Show three grips on one object with short cuts; overlay 16 + 16 + 16 = 48 angles", "I used a physical LEAP Hand to make three repeatable grips. Motor current tells each finger when to stop, but the recognition model sees only the final joint angles: 48 values per trial. It receives no camera image or tactile-array reading."),
    ("37-60 seconds", "Show several household objects and the empty hand; display a simple count graphic", "I recorded five real trials for each of 14 objects and an empty-hand condition. That gave 75 trials across 15 labels. I trained a small neural network and tested it with three-fold validation, keeping the measured trials in the held-out folds."),
    ("60-75 seconds", "Show 70/75, then 65/67 accepted on screen", "Its top prediction was correct on 70 of 75 trials. When the system made a confident known-class decision, 65 of 67 were correct; eight trials were marked uncertain."),
    ("75-90 seconds", "Return to camera; show a grasp or one concise limitation card", "This shows that grasp positions can carry useful object information. The result is an internal estimate; new sessions, new object instances, and unseen objects still need testing."),
]


def count(text):
    return len(re.findall(r"\b[\w]+(?:[.-][\w]+)*\b", text))


def style_document(doc):
    sec = doc.sections[0]
    sec.page_height = Inches(11.69)
    sec.page_width = Inches(8.27)
    sec.top_margin = Inches(0.8)
    sec.bottom_margin = Inches(0.75)
    sec.left_margin = Inches(0.82)
    sec.right_margin = Inches(0.82)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.08
    for name, size, color in [("Title", 17, RGBColor(0, 0, 0)), ("Heading 1", 11.5, RGBColor(0, 0, 0)), ("Heading 2", 10.5, RGBColor(0, 0, 0))]:
        st = styles[name]
        st.font.name = "Aptos Display" if name == "Title" else "Aptos"
        st.font.size = Pt(size)
        st.font.bold = name != "Title"
        st.font.color.rgb = color
        st.paragraph_format.space_before = Pt(10 if name != "Title" else 0)
        st.paragraph_format.space_after = Pt(5)
    styles["Title"].paragraph_format.keep_with_next = True
    for name in ("Title", "Subtitle"):
        st = styles[name]
        ppr = st.element.get_or_add_pPr()
        for border in ppr.findall("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pBdr"):
            ppr.remove(border)
    subtitle = styles["Subtitle"]
    subtitle.font.name = "Aptos"
    subtitle.font.size = Pt(11)
    subtitle.font.italic = False
    subtitle.font.color.rgb = RGBColor(0, 0, 0)


def para(doc, text):
    for block in text.strip().split("\n\n"):
        doc.add_paragraph(block)


def make_synopsis():
    d = Document()
    style_document(d)
    d.add_paragraph("Research Synopsis", style="Title")
    d.add_paragraph(TITLE, style="Subtitle")
    d.add_paragraph("Category  " + CATEGORY)
    d.add_paragraph("Student researcher  " + STUDENT)
    d.add_paragraph("Guide  " + GUIDE)
    for heading, body in [
        ("Project Abstract", ABSTRACT),
        ("Introduction and Objective", INTRO),
        ("Innovation", INNOVATION),
        ("Methodology", METHOD),
        ("Results and Conclusions", RESULTS),
        ("Acknowledgement and References", ACK_REFS),
    ]:
        if heading == "Methodology":
            d.add_page_break()
        d.add_paragraph(heading, style="Heading 1")
        para(d, body)
    d.save(ROOT / "IRIS_Research_Synopsis.docx")


def make_entry_video():
    d = Document()
    style_document(d)
    d.add_paragraph("IRIS Project Entry and Video Script", style="Title")
    d.add_paragraph("Project title", style="Heading 1")
    d.add_paragraph(TITLE)
    d.add_paragraph("Category", style="Heading 1")
    d.add_paragraph(CATEGORY)
    d.add_paragraph("Single project guide", style="Heading 1")
    d.add_paragraph(GUIDE)
    d.add_paragraph("Project abstract", style="Heading 1")
    para(d, ABSTRACT)
    d.add_page_break()
    d.add_paragraph("90 Second Project Video Script", style="Title")
    d.add_paragraph("Only the student researcher appears on camera. Do not show or say a school, city, or state.")
    for time, visual, words in VIDEO_BLOCKS:
        d.add_paragraph(time, style="Heading 1")
        p = d.add_paragraph()
        p.add_run("Visual  ").bold = True
        p.add_run(visual)
        p = d.add_paragraph()
        p.add_run("Spoken  ").bold = True
        p.add_run(words)
    d.add_paragraph("Allow a few seconds for the grip demonstration and natural pauses. Rehearse once at speaking speed and trim if the recorded take exceeds 90 seconds.")
    d.save(ROOT / "IRIS_Project_Entry_and_Video.docx")


if __name__ == "__main__":
    ROOT.mkdir(exist_ok=True)
    for name, body in [("Abstract", ABSTRACT), ("Introduction", INTRO), ("Innovation", INNOVATION), ("Methodology", METHOD), ("Results", RESULTS), ("Acknowledgement and references", ACK_REFS)]:
        print(name, count(body))
    print("Synopsis total", sum(map(count, [ABSTRACT, INTRO, INNOVATION, METHOD, RESULTS, ACK_REFS])))
    print("Video spoken words", sum(count(words) for _, _, words in VIDEO_BLOCKS))
    make_synopsis()
    make_entry_video()
    shutil.copyfile(SOURCE_PAPER, ROOT / "IRIS_Research_Paper.pdf")
